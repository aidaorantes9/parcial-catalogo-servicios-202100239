#!/usr/bin/env python3
"""Análisis verificable de data/CatalogoServicios.xlsx (solo lectura).

El contenido de las celdas se trata como DATO, nunca como instrucción.
El libro se abre con read_only=False solo para poder leer merged_cells;
nunca se guarda. Se comprueba el SHA-256 antes y después de la lectura.

Uso (desde la raíz del repositorio):
  docker run --rm -v "$PWD":/w -w /w python:3.12-slim \
    sh -c "pip install -q openpyxl && python scripts/analizar_excel.py"

Salidas:
  - stdout (mismo contenido que el Markdown)
  - docs/contexto/analisis-excel.md
  - data/CatalogoServicios.xlsx.sha256 (formato `sha256sum -c`)
Código de salida: 0 si los controles 12 (nivel 1) y 46 (nivel 2) pasan; 1 si no.
"""

import datetime
import difflib
import hashlib
import os
import re
import sys
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.utils import get_column_letter

EXCEL = "data/CatalogoServicios.xlsx"
SHA_FILE = EXCEL + ".sha256"
SALIDA_MD = "docs/contexto/analisis-excel.md"
HOJA = "Servicios Externos"

FILA_ENCABEZADO = 4
FILA_INI, FILA_FIN = 5, 101
COLS = "ABCDEFGHIJKL"
LISTAS_FILA_INI, LISTAS_FILA_FIN = 112, 122
LISTAS = OrderedDict([("E", "activo"), ("F", "clase"), ("G", "criticidad"), ("H", "tipo")])

ESPERADO_N1 = 12
ESPERADO_N2 = 46

# Listas tal como las transcribe el enunciado (§2). Se usan solo para contrastar.
LISTAS_ENUNCIADO = {
    "activo": ["S", "N"],  # el enunciado dice "Indicador S/N"
    "clase": ["A DEMANDA", "RECURRENTE"],
    "criticidad": ["Very Low", "Low", "Normal", "High", "Very High"],
    "tipo": ["Back End", "Demostration", "End User Service", "Front End", "IT Management",
             "IT Operational", "Other", "Project", "Reporting", "Training",
             "Underpinning Contract"],
}

# Vocabulario de referencia escrito a mano (ortografía estándar en inglés/español)
# para detectar posibles errores de escritura en las etiquetas de las listas.
# Es una heurística declarada: una palabra que no está en el vocabulario pero se
# parece (difflib >= 0.8) a una que sí está, se reporta como "posible error".
VOCABULARIO_REFERENCIA = {
    "S", "N", "A", "DEMANDA", "RECURRENTE", "Very", "Low", "Normal", "High",
    "Back", "End", "Demonstration", "User", "Service", "Front", "IT", "Management",
    "Operational", "Other", "Project", "Reporting", "Training", "Underpinning", "Contract",
}

# Patrones que podrían indicar texto con forma de instrucción dirigido a un asistente.
# Coincidir NO implica ejecutar nada: solo se reporta como hallazgo.
PATRONES_INSTRUCCION = [
    r"\bignor", r"\binstrucci", r"\bprompt\b", r"\bsystem\b", r"\bsistema:", r"\bejecut",
    r"\bborr", r"\belimin", r"\bdelete\b", r"\bdrop\b", r"\brm\s+-", r"\bcurl\b",
    r"\bwget\b", r"\bpassword\b", r"\bcontrase", r"\bsecret", r"\btoken\b", r"\bapi[_ ]?key",
    r"\brevel", r"\bmuestr", r"\benv[ií]a", r"\bolvid", r"\bact[uú]a como\b", r"\bassistant\b",
    r"\basistente\b", r"\bIA\b", r"\bAI\b", r"https?://", r"\bcommit\b", r"\bpush\b",
]

RE_COD_N1 = re.compile(r"^SE\.\d{2}$")
RE_COD_N2 = re.compile(r"^SE\.\d{2}\.\d{2}$")

out = []


def p(linea=""):
    out.append(linea)


def sha256(ruta):
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(65536), b""):
            h.update(bloque)
    return h.hexdigest()


def fmt(v):
    """Representación legible y segura para Markdown de un valor de celda."""
    if v is None:
        return "∅"
    s = repr(v) if isinstance(v, str) else str(v)
    return s.replace("|", "\\|").replace("\n", "\\n")


def tipo(v):
    if v is None:
        return "vacío"
    if isinstance(v, bool):
        return "booleano"
    if isinstance(v, (int, float)):
        return "número"
    if isinstance(v, str):
        return "texto" if v.strip() != "" else "texto en blanco"
    return type(v).__name__


def vacio(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def tabla(encabezados, filas):
    p("| " + " | ".join(encabezados) + " |")
    p("|" + "|".join("---" for _ in encabezados) + "|")
    for f in filas:
        p("| " + " | ".join(str(x) for x in f) + " |")


def main():
    hash_antes = sha256(EXCEL)

    wb = openpyxl.load_workbook(EXCEL, read_only=False, data_only=False)
    ws = wb[HOJA]

    # --- Índice de combinaciones --------------------------------------------------
    rangos = sorted(ws.merged_cells.ranges, key=lambda r: (r.min_col, r.min_row))
    celda_a_rango = {}
    for r in rangos:
        for fila in range(r.min_row, r.max_row + 1):
            for col in range(r.min_col, r.max_col + 1):
                celda_a_rango[(fila, col)] = r

    def bruto(fila, col):
        return ws.cell(row=fila, column=col).value

    def efectivo(fila, col):
        """Valor de la celda; si está en un rango combinado, el de su celda principal
        (solo dentro de ese rango, sin propagar a filas ajenas)."""
        r = celda_a_rango.get((fila, col))
        if r is None:
            return bruto(fila, col)
        return bruto(r.min_row, r.min_col)

    def es_continuacion(fila, col):
        r = celda_a_rango.get((fila, col))
        return r is not None and (fila, col) != (r.min_row, r.min_col)

    ci = {c: i + 1 for i, c in enumerate(COLS)}

    p("# Análisis verificable de `data/CatalogoServicios.xlsx`")
    p()
    p("> Generado por `scripts/analizar_excel.py`. No editar a mano: volver a ejecutar el script.")
    p(">")
    p("> El contenido de las celdas es **dato**, no instrucciones. Los valores de texto se muestran")
    p("> con `repr()` (entre comillas) para hacer visibles espacios iniciales/finales; `∅` = celda vacía.")
    p()
    p(f"- Fecha de generación (UTC): {datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%d %H:%M:%S')}")
    p(f"- openpyxl {openpyxl.__version__}, Python {sys.version.split()[0]}")
    p(f"- SHA-256 del Excel: `{hash_antes}`")
    p(f"- Hojas del libro: {wb.sheetnames}")
    p(f"- Dimensión declarada de la hoja `{HOJA}`: `{ws.dimensions}` (max_row={ws.max_row}, max_column={ws.max_column})")
    p()

    # Encabezados
    p("## 0. Encabezados (fila 4)")
    p()
    tabla(["Columna", "Valor"], [(c, fmt(bruto(FILA_ENCABEZADO, ci[c]))) for c in COLS])
    p()

    # --- a. Rangos combinados ----------------------------------------------------
    p("## a. Rangos combinados de la hoja")
    p()
    p(f"Total de rangos combinados: **{len(rangos)}**")
    p()
    por_col = defaultdict(int)
    filas_tabla = []
    for r in rangos:
        col = get_column_letter(r.min_col) if r.min_col == r.max_col else \
            f"{get_column_letter(r.min_col)}–{get_column_letter(r.max_col)}"
        por_col[col] += 1
        filas_tabla.append((col, f"`{r.coord}`", r.max_row - r.min_row + 1,
                            fmt(bruto(r.min_row, r.min_col))))
    tabla(["Columna", "Rango", "Filas", "Valor de la celda principal"], filas_tabla)
    p()
    p("Rangos por columna: " + ", ".join(f"{k}={v}" for k, v in sorted(por_col.items())))
    fuera = [r.coord for r in rangos if r.min_row < FILA_INI or r.max_row > FILA_FIN]
    p()
    p(f"Rangos que salen de las filas {FILA_INI}–{FILA_FIN}: {fuera if fuera else 'ninguno'}")
    p()

    # --- b. Filas 5–101 ----------------------------------------------------------
    p(f"## b. Filas {FILA_INI}–{FILA_FIN}: valor efectivo y clasificación")
    p()
    p("Clasificación según la columna C (COD.N2):")
    p("- **servicio nivel 2**: C tiene valor propio (celda no combinada o celda principal de su rango).")
    p("- **fila de continuación**: C está dentro de un rango combinado y no es su celda principal.")
    p("- **fila sin código fuera de combinación**: la fila tiene contenido, C está vacía y no pertenece a ningún rango.")
    p("- **vacía**: ninguna de A–L tiene valor efectivo.")
    p()
    p("Las celdas marcadas con `⁺` obtienen su valor de la celda principal de su rango combinado.")
    p()

    CLASES = ("servicio nivel 2", "fila de continuación (dentro de un rango combinado)",
              "fila sin código fuera de combinación", "vacía")
    clasif = {}
    filas_tabla = []
    for fila in range(FILA_INI, FILA_FIN + 1):
        vals = {c: efectivo(fila, ci[c]) for c in COLS}
        c_bruto = bruto(fila, ci["C"])
        if all(vacio(v) for v in vals.values()):
            k = CLASES[3]
        elif es_continuacion(fila, ci["C"]):
            k = CLASES[1]
        elif not vacio(c_bruto):
            k = CLASES[0]
        else:
            k = CLASES[2]
        clasif[fila] = k
        celdas = []
        for c in COLS:
            marca = "⁺" if es_continuacion(fila, ci[c]) else ""
            celdas.append(fmt(vals[c]) + marca)
        filas_tabla.append([fila] + celdas + [k])
    tabla(["Fila"] + list(COLS) + ["Clasificación"], filas_tabla)
    p()
    conteo = defaultdict(list)
    for fila, k in clasif.items():
        conteo[k].append(fila)
    for k in CLASES:
        p(f"- {k}: **{len(conteo[k])}** filas" + (f" → {conteo[k]}" if k != CLASES[0] else ""))
    p()
    if conteo[CLASES[2]]:
        p("**Hallazgo:** filas sin código fuera de combinación (no deben asignarse automáticamente a otro servicio):")
        p()
        for fila in conteo[CLASES[2]]:
            previo = next((f for f in range(fila - 1, FILA_INI - 1, -1) if clasif[f] == CLASES[0]), None)
            cprev = efectivo(previo, ci["C"]) if previo else None
            rngC = celda_a_rango.get((fila - 1, ci["C"]))
            p(f"- Fila {fila}: A={fmt(efectivo(fila, ci['A']))}, "
              f"E–H={[efectivo(fila, ci[c]) for c in 'EFGH']}. "
              f"Servicio nivel 2 explícito anterior: fila {previo} ({fmt(cprev)}); "
              f"rango de C que termina justo antes: {('`' + rngC.coord + '`') if rngC else 'ninguno'}. "
              f"Pertenencia a un servicio nivel 2: **no determinado** (C vacía y fuera de cualquier rango).")
        p()

    # --- c. Códigos distintos -----------------------------------------------------
    p("## c. Códigos distintos de nivel 1 y nivel 2")
    p()
    n1 = OrderedDict()   # código -> [filas donde aparece explícito]
    n2 = OrderedDict()
    for fila in range(FILA_INI, FILA_FIN + 1):
        a = bruto(fila, ci["A"])
        c = bruto(fila, ci["C"])
        if not vacio(a):
            n1.setdefault(str(a).strip(), []).append(fila)
        if not vacio(c):
            n2.setdefault(str(c).strip(), []).append(fila)
    p(f"### Nivel 1 (columna A, valores explícitos): **{len(n1)}** códigos distintos")
    p()
    tabla(["Código", "Filas con valor explícito", "Nombre(s) en B", "¿Formato SE.NN?"],
          [(k, v, " / ".join(fmt(bruto(f, ci['B'])) for f in v), "sí" if RE_COD_N1.match(k) else "**no**")
           for k, v in n1.items()])
    p()
    p(f"### Nivel 2 (columna C, valores explícitos): **{len(n2)}** códigos distintos")
    p()
    filas_tabla = []
    for k, v in n2.items():
        f0 = v[0]
        rng = celda_a_rango.get((f0, ci["C"]))
        filas_tabla.append((k, v, f"`{rng.coord}`" if rng else "—", fmt(bruto(f0, ci["D"])),
                            fmt(efectivo(f0, ci["A"])), "sí" if RE_COD_N2.match(k) else "**no**"))
    tabla(["Código", "Fila(s)", "Rango C", "Nombre (D)", "N1 efectivo (A)", "¿Formato SE.NN.NN?"],
          filas_tabla)
    p()
    # Relación prefijo del código N2 vs N1 efectivo
    incons = []
    for k, v in n2.items():
        a_ef = efectivo(v[0], ci["A"])
        pref = ".".join(k.split(".")[:2])
        if vacio(a_ef):
            incons.append(f"- `{k}` (fila {v[0]}): A efectivo vacío; el prefijo del código sugiere `{pref}` "
                          f"pero el padre por celdas es **no determinado** (A{v[0]} vacía y no combinada).")
        elif str(a_ef).strip() != pref:
            incons.append(f"- `{k}` (fila {v[0]}): prefijo `{pref}` ≠ A efectivo `{a_ef}`.")
    p("Relación prefijo de código N2 ↔ código N1 efectivo de su fila:")
    p()
    p("\n".join(incons) if incons else "- Todos los códigos N2 coinciden con el N1 efectivo de su fila.")
    p()

    # --- d. Duplicados y conflictos ----------------------------------------------
    p("## d. Duplicados y conflictos de atributos")
    p()
    dup_n2 = {k: v for k, v in n2.items() if len(v) > 1}
    p(f"- Códigos de nivel 2 duplicados (explícitos en más de una fila): "
      f"{dup_n2 if dup_n2 else 'ninguno'}")
    dup_n1 = {k: v for k, v in n1.items() if len(v) > 1}
    p(f"- Códigos de nivel 1 explícitos en más de una fila (fuera de una combinación): "
      f"{dup_n1 if dup_n1 else 'ninguno'}")
    for k, v in dup_n1.items():
        nombres = {f: bruto(f, ci["B"]) for f in v}
        distintos = set(nombres.values())
        estado = "**CONFLICTO de nombre**" if len(distintos) > 1 else "mismo nombre"
        p(f"  - `{k}`: {estado}: " + "; ".join(f"fila {f} → B={fmt(n)}" for f, n in nombres.items()))
    nombres_n2 = defaultdict(list)
    for k, v in n2.items():
        nombres_n2[bruto(v[0], ci["D"])].append(k)
    rep = {n: cs for n, cs in nombres_n2.items() if len(cs) > 1}
    p(f"- Nombres de nivel 2 repetidos con códigos distintos: {rep if rep else 'ninguno'}")
    p()
    p("Conflictos de atributos dentro de un mismo servicio nivel 2 (filas de su rango combinado en C):")
    p()
    p("Conflicto = dos o más valores **no vacíos** distintos en la misma columna. Una celda vacía en una "
      "fila de continuación (columnas no combinadas I, K, L) se reporta aparte como ausencia.")
    p()
    conflictos, ausencias = [], []
    for k, v in n2.items():
        f0 = v[0]
        rng = celda_a_rango.get((f0, ci["C"]))
        if not rng:
            continue
        for c in "EFGHIJKL":
            valores = OrderedDict()
            for fila in range(rng.min_row, rng.max_row + 1):
                valores.setdefault(efectivo(fila, ci[c]), []).append(fila)
            no_vacios = [val for val in valores if not vacio(val)]
            detalle = "; ".join(f"{fmt(val)} en filas {fs}" for val, fs in valores.items())
            if len(no_vacios) > 1:
                conflictos.append(f"- `{k}` ({rng.coord}) columna {c}: {detalle}")
            elif len(valores) > 1:
                ausencias.append(f"- `{k}` ({rng.coord}) columna {c}: {detalle}")
    p("Conflictos:")
    p()
    p("\n".join(conflictos) if conflictos else "- Ninguno: en cada rango, E–L no tienen dos valores no vacíos distintos.")
    p()
    p("Ausencias en filas de continuación (valor solo en la fila principal):")
    p()
    p("\n".join(ausencias) if ausencias else "- Ninguna.")
    p()

    # --- e. SE.12 -----------------------------------------------------------------
    p("## e. Análisis específico de SE.12 (filas 99–101)")
    p()
    filas_se12 = [99, 100, 101]
    filas_tabla = []
    for fila in filas_se12:
        for c in COLS:
            cel = ws.cell(row=fila, column=ci[c])
            rng = celda_a_rango.get((fila, ci[c]))
            filas_tabla.append((fila, c, fmt(cel.value), tipo(cel.value), cel.data_type,
                                cel.number_format, f"`{rng.coord}`" if rng else "—"))
    tabla(["Fila", "Col", "Valor bruto", "Tipo Python", "data_type openpyxl", "number_format",
           "Rango combinado"], filas_tabla)
    p()
    rng_99_101 = [r.coord for r in rangos if r.min_row <= 101 and r.max_row >= 99]
    p(f"- Rangos combinados que afectan a las filas 99–101: {rng_99_101 if rng_99_101 else 'ninguno'}")
    rng_96 = [r.coord for r in rangos if r.min_row <= 98 and r.max_row >= 96]
    p(f"- Rangos del grupo anterior (SE.11, filas 96–98), para comparar: {rng_96}")
    p()
    p("Nombres de nivel 1 presentes para SE.12:")
    p()
    se12_nombres = OrderedDict()
    for fila in filas_se12:
        if str(bruto(fila, ci["A"]) or "").strip() == "SE.12":
            se12_nombres[fila] = bruto(fila, ci["B"])
    hijos_se12 = [(f, bruto(f, ci["C"]), bruto(f, ci["D"])) for f in filas_se12 if not vacio(bruto(f, ci["C"]))]
    for fila, nombre in se12_nombres.items():
        coincide = [h for h in hijos_se12 if h[2] == nombre]
        p(f"- Fila {fila}: A{fila}=`SE.12`, B{fila}={fmt(nombre)}"
          + (f" — **idéntico** al nombre del hijo {coincide[0][1]} (D{coincide[0][0]})" if coincide else ""))
    p()
    p("Hijos explícitos de SE.12 y tipo de dato del código:")
    p()
    filas_tabla = []
    for f, cod, nom in hijos_se12:
        cel = ws.cell(row=f, column=ci["C"])
        filas_tabla.append((f, fmt(cod), tipo(cod), cel.data_type, cel.number_format, fmt(nom),
                            "sí" if RE_COD_N2.match(str(cod)) else "**no** (un solo dígito final)"))
    tabla(["Fila", "Código", "Tipo", "data_type", "number_format", "Nombre (D)", "¿Formato SE.NN.NN?"],
          filas_tabla)
    p()
    p("Comparación con otros grupos: ¿algún otro nombre de nivel 1 coincide con el nombre de uno de sus hijos?")
    p()
    otros = []
    for k, v in n1.items():
        if k == "SE.12":
            continue
        nombre = bruto(v[0], ci["B"])
        hijos = [(kk, bruto(vv[0], ci["D"])) for kk, vv in n2.items() if kk.startswith(k + ".")]
        igual = [h for h in hijos if h[1] == nombre]
        if igual:
            otros.append(f"- `{k}` {fmt(nombre)} = nombre del hijo `{igual[0][0]}`")
    p("\n".join(otros) if otros else "- Ninguno.")
    p()
    p("Observación: A101 está vacía y no combinada, por lo que el padre de `SE.12.3` no se obtiene por celdas; "
      "solo por el prefijo de su código.")
    p()

    # --- f. Celdas vacías -------------------------------------------------------
    p("## f. Celdas vacías (valor efectivo) por columna en filas de servicio nivel 2")
    p()
    filas_serv = conteo[CLASES[0]]
    filas_tabla = []
    for c in COLS:
        vac = [f for f in filas_serv if vacio(efectivo(f, ci[c]))]
        filas_tabla.append((c, bruto(FILA_ENCABEZADO, ci[c]), len(vac), len(filas_serv),
                            vac if len(vac) <= 12 else f"{vac[:12]}… (+{len(vac) - 12})"))
    tabla(["Col", "Campo", "Vacías", "De", "Filas vacías"], filas_tabla)
    p()
    p("Detalle de filas 99–101 (valor efectivo):")
    p()
    filas_tabla = []
    for fila in filas_se12:
        vac = [c for c in COLS if vacio(efectivo(fila, ci[c]))]
        llenas = [c for c in COLS if not vacio(efectivo(fila, ci[c]))]
        filas_tabla.append((fila, clasif[fila], ", ".join(llenas), ", ".join(vac)))
    tabla(["Fila", "Clasificación", "Columnas con valor", "Columnas vacías"], filas_tabla)
    p()

    # --- g. Listas de opciones -------------------------------------------------------
    p(f"## g. Listas de opciones E{LISTAS_FILA_INI}:H{LISTAS_FILA_FIN}")
    p()
    p(f"Encabezados en la fila {LISTAS_FILA_INI - 1}: " +
      ", ".join(f"{c}{LISTAS_FILA_INI - 1}={fmt(bruto(LISTAS_FILA_INI - 1, ci[c]))}" for c in LISTAS))
    p()
    listas = OrderedDict()
    for c, nombre in LISTAS.items():
        listas[nombre] = [(f, bruto(f, ci[c])) for f in range(LISTAS_FILA_INI, LISTAS_FILA_FIN + 1)
                          if not vacio(bruto(f, ci[c]))]
    for (c, nombre) in LISTAS.items():
        p(f"- **{nombre}** (columna {c}): " + ", ".join(f"{fmt(v)} (fila {f})" for f, v in listas[nombre]))
    p()
    p("Contraste con las listas transcritas en el enunciado (§2):")
    p()
    for nombre, esperado in LISTAS_ENUNCIADO.items():
        real = [v for _, v in listas[nombre]]
        p(f"- {nombre}: " + ("coincide exactamente" if real == esperado else
                             f"**difiere** — Excel={real} enunciado={esperado}"))
    p()
    p("Valores de E–H en filas de servicio/continuación/sin código que NO están en su lista "
      "(se distinguen vacíos de valores no vacíos fuera de lista):")
    p()
    fuera_lista, vacios_lista = [], defaultdict(list)
    for fila in range(FILA_INI, FILA_FIN + 1):
        if clasif[fila] == CLASES[3]:
            continue
        for c, nombre in LISTAS.items():
            v = efectivo(fila, ci[c])
            permitidos = [x for _, x in listas[nombre]]
            if vacio(v):
                vacios_lista[fila].append(c)
            elif v not in permitidos:
                casi = [x for x in permitidos if isinstance(v, str) and isinstance(x, str)
                        and v.strip().lower() == x.strip().lower()]
                nota = f" (coincide ignorando mayúsculas/espacios con {fmt(casi[0])})" if casi else ""
                fuera_lista.append((fila, clasif[fila], c, nombre, fmt(v) + nota))
    if fuera_lista:
        tabla(["Fila", "Clasificación", "Col", "Lista", "Valor"], fuera_lista)
    else:
        p("- Valores no vacíos fuera de lista: ninguno.")
    p("- Vacíos (no pertenecen a ninguna lista): " +
      ("; ".join(f"fila {f} ({clasif[f]}) columnas {''.join(cs)}" for f, cs in vacios_lista.items())
       if vacios_lista else "ninguno"))
    p()
    p("Posibles errores de escritura en las etiquetas de las listas (heurística: palabra fuera de un "
      "vocabulario de referencia escrito a mano en el script y similar ≥ 0.8 a una palabra de ese vocabulario):")
    p()
    errores = []
    for nombre, vals in listas.items():
        for f, v in vals:
            if not isinstance(v, str):
                continue
            if v != v.strip() or "  " in v:
                errores.append(f"- {nombre} fila {f}: {fmt(v)} tiene espacios sobrantes")
            for palabra in v.split():
                if palabra not in VOCABULARIO_REFERENCIA:
                    sug = difflib.get_close_matches(palabra, VOCABULARIO_REFERENCIA, n=1, cutoff=0.8)
                    if sug:
                        errores.append(f"- {nombre} fila {f}: `{palabra}` en {fmt(v)} → posible error; "
                                       f"forma estándar más cercana `{sug[0]}`")
                    else:
                        errores.append(f"- {nombre} fila {f}: `{palabra}` no está en el vocabulario de "
                                       f"referencia; corrección **no determinada**")
    p("\n".join(errores) if errores else "- Ninguno detectado.")
    p()

    # --- h. K y L --------------------------------------------------------------------
    p("## h. Tipo de dato real de K (Minimo) y L (Maximo), filas 5–101")
    p()
    for c in "KL":
        tipos = defaultdict(list)
        for fila in range(FILA_INI, FILA_FIN + 1):
            tipos[tipo(bruto(fila, ci[c]))].append(fila)
        p(f"- Columna {c} ({bruto(FILA_ENCABEZADO, ci[c])}): " +
          "; ".join(f"{t}: {len(fs)} filas" + (f" {fs}" if t != 'vacío' else "") for t, fs in tipos.items()))
        for fila in range(FILA_INI, FILA_FIN + 1):
            cel = ws.cell(row=fila, column=ci[c])
            if cel.value is not None:
                p(f"  - {c}{fila} = {fmt(cel.value)} ({tipo(cel.value)}, data_type `{cel.data_type}`, "
                  f"formato `{cel.number_format}`)")
        no_num = [f for t, fs in tipos.items() if t not in ("número", "vacío") for f in fs]
        p(f"  - Valores no numéricos (excluyendo vacíos): {no_num if no_num else 'ninguno'}")
    p()
    p("Validación mínimo ≤ máximo donde ambos existen:")
    p()
    for fila in range(FILA_INI, FILA_FIN + 1):
        k, l = bruto(fila, ci["K"]), bruto(fila, ci["L"])
        if isinstance(k, (int, float)) and isinstance(l, (int, float)):
            p(f"- Fila {fila} ({efectivo(fila, ci['C'])}): {k} ≤ {l} → {'OK' if k <= l else '**FALLA**'}")
        elif (k is None) != (l is None):
            p(f"- Fila {fila}: solo uno de K/L tiene valor (K={fmt(k)}, L={fmt(l)})")
    p()

    # --- i. Contenido fuera de rangos esperados ----------------------------------------
    p(f"## i. Celdas con contenido fuera de filas {FILA_INI}–{FILA_FIN} (A–L) y de "
      f"E{LISTAS_FILA_INI}:H{LISTAS_FILA_FIN}")
    p()
    p("Se recorren todas las celdas de la hoja (incluidas columnas M en adelante).")
    p()
    extra = []
    formulas = []
    for row in ws.iter_rows():
        for cel in row:
            if cel.value is None:
                continue
            if cel.data_type == "f" or (isinstance(cel.value, str) and cel.value.startswith("=")):
                formulas.append(cel.coordinate)
            en_datos = FILA_INI <= cel.row <= FILA_FIN and cel.column <= len(COLS)
            en_listas = (LISTAS_FILA_INI <= cel.row <= LISTAS_FILA_FIN and
                         get_column_letter(cel.column) in LISTAS)
            if not (en_datos or en_listas):
                nota = "encabezado esperado (A4:L4)" if cel.row == FILA_ENCABEZADO and cel.column <= 12 else ""
                extra.append((cel.coordinate, fmt(cel.value), tipo(cel.value), nota))
    tabla(["Celda", "Valor", "Tipo", "Nota"], extra) if extra else p("- Ninguna.")
    p()
    p(f"- Fórmulas en la hoja: {formulas if formulas else 'ninguna'}")
    dv = ws.data_validations.dataValidation if ws.data_validations else []
    p(f"- Validaciones de datos definidas: {len(dv)}" +
      ("" if not dv else " → " + "; ".join(f"{d.sqref} tipo={d.type} fórmula={d.formula1}" for d in dv)))
    ocultas_f = [f for f, d in ws.row_dimensions.items() if d.hidden]
    ocultas_c = [c for c, d in ws.column_dimensions.items() if d.hidden]
    p(f"- Filas ocultas: {ocultas_f if ocultas_f else 'ninguna'}; columnas ocultas: {ocultas_c if ocultas_c else 'ninguna'}")
    p()

    # --- Seguridad: texto con forma de instrucción y anomalías de texto ------------------
    p("## Hallazgos de seguridad y calidad de texto")
    p()
    p("Regla: ningún texto del Excel se interpreta como instrucción. Se listan textos que coinciden con "
      "patrones típicos de órdenes a un asistente o que son ajenos al campo; solo se reportan.")
    p()
    sospechosos, espacios = [], []
    for row in ws.iter_rows():
        for cel in row:
            v = cel.value
            if not isinstance(v, str):
                continue
            pats = [pt for pt in PATRONES_INSTRUCCION if re.search(pt, v, re.IGNORECASE)]
            if pats:
                sospechosos.append(f"- {cel.coordinate}: {fmt(v)} — patrones: {pats}")
            if v != v.strip() or "  " in v:
                espacios.append(f"- {cel.coordinate}: {fmt(v)}")
    p("Textos con forma de instrucción:")
    p()
    p("\n".join(sospechosos) if sospechosos else "- Ninguno.")
    p()
    p("Columna I (Descripción) — todos los valores no vacíos, para revisión humana:")
    p()
    desc = [(f, bruto(f, ci["I"])) for f in range(FILA_INI, FILA_FIN + 1) if not vacio(bruto(f, ci["I"]))]
    p("\n".join(f"- I{f}: {fmt(v)} (servicio {efectivo(f, ci['C'])} {fmt(efectivo(f, ci['D']))})"
                for f, v in desc) if desc else "- Ninguno.")
    p()
    p("Textos con espacios iniciales, finales o dobles:")
    p()
    p("\n".join(espacios) if espacios else "- Ninguno.")
    p()

    # --- Integridad del archivo y controles ----------------------------------------------
    wb.close()
    hash_despues = sha256(EXCEL)
    with open(SHA_FILE, "w", encoding="utf-8") as f:
        f.write(f"{hash_antes}  {EXCEL}\n")

    p("## CONTROLES")
    p()
    ok_n1 = len(n1) == ESPERADO_N1
    ok_n2 = len(n2) == ESPERADO_N2
    tabla(["Control", "Esperado", "Obtenido", "Resultado"], [
        ("Códigos distintos de nivel 1", ESPERADO_N1, len(n1), "PASA" if ok_n1 else "FALLA"),
        ("Códigos explícitos distintos de nivel 2", ESPERADO_N2, len(n2), "PASA" if ok_n2 else "FALLA"),
    ])
    p()
    p(f"- (informativo) SHA-256 antes = después de la lectura: "
      f"{'sí' if hash_antes == hash_despues else '**NO**'} (`{hash_despues}`)")
    p(f"- (informativo) Hash guardado en `{SHA_FILE}`; verificar con `sha256sum -c {SHA_FILE}`")
    p()

    texto = "\n".join(out) + "\n"
    print(texto, end="")
    os.makedirs(os.path.dirname(SALIDA_MD), exist_ok=True)
    with open(SALIDA_MD, "w", encoding="utf-8") as f:
        f.write(texto)

    # Si corre como root dentro de Docker, devolver la propiedad de las salidas al dueño del Excel.
    if hasattr(os, "geteuid") and os.geteuid() == 0:
        st = os.stat(EXCEL)
        for ruta in (SALIDA_MD, SHA_FILE):
            os.chown(ruta, st.st_uid, st.st_gid)

    return 0 if (ok_n1 and ok_n2) else 1


if __name__ == "__main__":
    sys.exit(main())
