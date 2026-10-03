"""Lectura del libro de Excel del catálogo (solo lectura, sin base de datos).

Adapta la lógica de `scripts/analizar_excel.py`. El contenido de las celdas es DATO: nunca se
interpreta como instrucción (AGENTS.md §9).

- El libro se abre desde los mismos bytes cuyo SHA-256 se verificó (BytesIO), así lo que se lee es
  exactamente lo verificado y el archivo en disco nunca se abre para escritura ni se guarda.
- Celdas combinadas: el valor de la celda principal se usa solo para las filas dentro de su rango;
  nunca se propaga fuera de él (enunciado §3.4.1).
- Clasificación de las filas 5–101 por la columna C: servicio de nivel 2 (C con valor propio), fila
  de continuación (C dentro de un rango y no es la principal), fila sin código fuera de combinación
  (D2) o fila vacía.
"""

import io
import re
from dataclasses import dataclass, field

import openpyxl
from openpyxl.utils import get_column_letter

HOJA = "Servicios Externos"
FILA_ENCABEZADO = 4
FILA_INI, FILA_FIN = 5, 101
COLUMNAS = "ABCDEFGHIJKL"
ENCABEZADOS = {
    "A": "COD.N1",
    "B": "SERVICIO - Nivel 1",
    "C": "COD.N2",
    "D": "SERVICIO - Nivel 2",
    "E": "ACTIVO",
    "F": "CLASE DE SERVICIO",
    "G": "CRITICIDAD",
    "H": "TIPO DE SERVICIO",
    "I": "Descripción",
    "J": "Métrica",
    "K": "Minimo",
    "L": "Maximo",
}
# Listas de opciones (filas ocultas). E (S/N) no es tabla: son dos valores fijos (modelo §3.2).
LISTAS_FILA_INI, LISTAS_FILA_FIN = 112, 122
LISTAS = {"F": "clase", "G": "criticidad", "H": "tipo"}

INDICE = {letra: i + 1 for i, letra in enumerate(COLUMNAS)}


class ErrorLectura(Exception):
    """El libro no tiene la estructura esperada; la importación no puede continuar."""


def vacio(valor):
    """Celda sin valor: None o cadena vacía. Un texto de solo espacios NO es vacío aquí: se
    conserva tal cual donde la regla lo exige (D8) y cada campo decide cómo tratarlo."""
    return valor is None or valor == ""


@dataclass
class Celda:
    celda: str
    valor: object

    def como_dict(self):
        return {"celda": self.celda, "valor": self.valor}


@dataclass
class Nivel1Leido:
    codigo: str
    filas: list
    rangos: list
    # Celdas B con valor propio (principal o no combinada) de las filas del código, en orden.
    nombres: list
    celda_codigo: Celda


@dataclass
class ServicioLeido:
    codigo: str
    fila: int
    filas: list
    rangos: list
    # Valor efectivo A–L en la fila principal; la celda indica de dónde sale el valor.
    valores: dict
    # Columnas cuyo valor viene de una celda principal de otra fila: {"A": "A5:A9"}.
    combinadas: dict
    # Valores de columnas no combinadas en filas de continuación: {"I": {6: None, 7: None}}.
    continuacion: dict = field(default_factory=dict)


@dataclass
class FilaOmitida:
    fila: int
    motivo: str  # "continuacion", "sin_codigo" o "vacia"
    valores: dict  # valor efectivo A–L


@dataclass
class ValorLista:
    columna: str
    fila: int
    etiqueta: object


@dataclass
class LibroLeido:
    nivel1: list
    nivel2: list
    omitidas: list
    listas: dict  # {"clase": [ValorLista, ...], ...}


def leer_libro(contenido):
    """Lee los bytes del libro y devuelve su estructura. No guarda nada."""
    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=False, data_only=False)
    except Exception as error:  # openpyxl lanza varios tipos según el daño del archivo
        raise ErrorLectura(f"No se pudo abrir el libro: {error}") from error
    try:
        if HOJA not in libro.sheetnames:
            raise ErrorLectura(f"El libro no tiene la hoja «{HOJA}» (hojas: {libro.sheetnames}).")
        return _Hoja(libro[HOJA]).leer()
    finally:
        libro.close()


class _Hoja:
    def __init__(self, hoja):
        self.hoja = hoja
        self.rango_de = {}
        for rango in hoja.merged_cells.ranges:
            for fila in range(rango.min_row, rango.max_row + 1):
                for columna in range(rango.min_col, rango.max_col + 1):
                    self.rango_de[(fila, columna)] = rango

    # ------------------------------------------------------------------ celdas

    def bruto(self, fila, letra):
        return self.hoja.cell(row=fila, column=INDICE[letra]).value

    def rango(self, fila, letra):
        return self.rango_de.get((fila, INDICE[letra]))

    def efectiva(self, fila, letra):
        """Celda de la que sale el valor: la principal de su rango (solo dentro del rango) o la
        propia celda."""
        rango = self.rango(fila, letra)
        if rango is None:
            return Celda(f"{letra}{fila}", self.bruto(fila, letra))
        principal = f"{get_column_letter(rango.min_col)}{rango.min_row}"
        return Celda(principal, self.hoja.cell(row=rango.min_row, column=rango.min_col).value)

    def es_continuacion(self, fila, letra):
        rango = self.rango(fila, letra)
        return rango is not None and (rango.min_row, rango.min_col) != (fila, INDICE[letra])

    # ------------------------------------------------------------------ lectura

    def leer(self):
        self._verificar_encabezados()
        nivel2, omitidas = [], []
        codigos_n2 = {}
        for fila in range(FILA_INI, FILA_FIN + 1):
            valores = {letra: self.efectiva(fila, letra).valor for letra in COLUMNAS}
            if all(vacio(v) for v in valores.values()):
                omitidas.append(FilaOmitida(fila, "vacia", valores))
            elif self.es_continuacion(fila, "C"):
                omitidas.append(FilaOmitida(fila, "continuacion", valores))
            elif not vacio(self.bruto(fila, "C")):
                servicio = self._servicio(fila)
                if servicio.codigo in codigos_n2:
                    raise ErrorLectura(
                        f"El código de nivel 2 «{servicio.codigo}» aparece en las filas "
                        f"{codigos_n2[servicio.codigo]} y {fila}; no se importa un código "
                        "duplicado."
                    )
                codigos_n2[servicio.codigo] = fila
                nivel2.append(servicio)
            else:
                omitidas.append(FilaOmitida(fila, "sin_codigo", valores))
        return LibroLeido(
            nivel1=self._nivel1(), nivel2=nivel2, omitidas=omitidas, listas=self._listas()
        )

    def _verificar_encabezados(self):
        distintos = [
            f"{letra}{FILA_ENCABEZADO}={self.bruto(FILA_ENCABEZADO, letra)!r} "
            f"(se esperaba {esperado!r})"
            for letra, esperado in ENCABEZADOS.items()
            if self.bruto(FILA_ENCABEZADO, letra) != esperado
        ]
        if distintos:
            raise ErrorLectura("Encabezados inesperados: " + "; ".join(distintos))

    def _servicio(self, fila):
        rango_c = self.rango(fila, "C")
        filas = list(range(rango_c.min_row, rango_c.max_row + 1)) if rango_c else [fila]
        valores, combinadas, rangos = {}, {}, []
        for letra in COLUMNAS:
            valores[letra] = self.efectiva(fila, letra)
            rango = self.rango(fila, letra)
            if rango is None:
                continue
            if self.es_continuacion(fila, letra):
                combinadas[letra] = rango.coord
            # A y B son rangos del nivel 1; el origen del nivel 2 guarda los de sus columnas.
            if letra not in "AB":
                rangos.append(rango.coord)
        continuacion = {}
        for letra in COLUMNAS[4:]:  # E–L
            if any(self.rango(f, letra) for f in filas):
                continue  # columna combinada: su valor ya es el de la celda principal
            otras = {f: self.bruto(f, letra) for f in filas if f != fila}
            if otras:
                continuacion[letra] = otras
        return ServicioLeido(
            codigo=_texto_codigo(self.bruto(fila, "C")),
            fila=fila,
            filas=filas,
            rangos=rangos,
            valores=valores,
            combinadas=combinadas,
            continuacion=continuacion,
        )

    def _nivel1(self):
        """Un registro por código explícito de A; sus filas son las que lo tienen como valor
        efectivo (dentro de su rango, sin propagar)."""
        por_codigo = {}
        for fila in range(FILA_INI, FILA_FIN + 1):
            if self.es_continuacion(fila, "A") or vacio(self.bruto(fila, "A")):
                continue
            codigo = _texto_codigo(self.bruto(fila, "A"))
            if codigo not in por_codigo:
                por_codigo[codigo] = Nivel1Leido(
                    codigo=codigo,
                    filas=[],
                    rangos=[],
                    nombres=[],
                    celda_codigo=Celda(f"A{fila}", self.bruto(fila, "A")),
                )
        for fila in range(FILA_INI, FILA_FIN + 1):
            valor_a = self.efectiva(fila, "A").valor
            if vacio(valor_a):
                continue
            n1 = por_codigo[_texto_codigo(valor_a)]
            n1.filas.append(fila)
            for letra in "AB":
                rango = self.rango(fila, letra)
                if rango is not None and rango.coord not in n1.rangos:
                    n1.rangos.append(rango.coord)
            if not self.es_continuacion(fila, "B") and not vacio(self.bruto(fila, "B")):
                n1.nombres.append(Celda(f"B{fila}", self.bruto(fila, "B")))
        return list(por_codigo.values())

    def _listas(self):
        return {
            nombre: [
                ValorLista(letra, fila, self.bruto(fila, letra))
                for fila in range(LISTAS_FILA_INI, LISTAS_FILA_FIN + 1)
                if not vacio(self.bruto(fila, letra))
            ]
            for letra, nombre in LISTAS.items()
        }


def _texto_codigo(valor):
    """Los códigos se conservan como texto exacto (D5). Un número se convierte a su texto."""
    return valor if isinstance(valor, str) else str(valor)


RE_COD_N1 = re.compile(r"^SE\.\d{2}$")
RE_COD_N2 = re.compile(r"^SE\.\d{2}\.\d{2}$")
